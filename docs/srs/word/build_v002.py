"""Build PBL6-SRS-v002.docx from v001 without overwriting the previous Word release."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "docs/srs/word/releases/PBL6-SRS-v001.docx"
DST = ROOT / "docs/srs/word/releases/PBL6-SRS-v002.docx"


def nb(s: str) -> str:
    return s.replace(" ", "\xa0")


def cell_text(cell) -> str:
    return cell.text.replace("\xa0", " ").replace("\u00a0", " ").strip()


def set_cell(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.text = text
        return
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(text)
    for extra in paragraphs[1:]:
        parent = extra._element.getparent()
        if parent is not None:
            parent.remove(extra._element)


def replace_in_cell(cell, old: str, new: str) -> bool:
    current = cell.text
    variants = {old, old.replace(" ", "\xa0"), nb(old)}
    updated = current
    for variant in variants:
        if variant in updated:
            updated = updated.replace(variant, new)
    if updated == current:
        plain = current.replace("\xa0", " ")
        target = old.replace("\xa0", " ")
        if target in plain:
            updated = plain.replace(target, new)
    if updated != current:
        set_cell(cell, updated)
        return True
    return False


def replace_all_cells(doc: Document, old: str, new: str) -> int:
    count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if replace_in_cell(cell, old, new):
                    count += 1
    return count


def replace_all_paragraphs(doc: Document, old: str, new: str) -> int:
    count = 0
    for para in doc.paragraphs:
        text = para.text
        if not text:
            continue
        variants = {old, old.replace(" ", "\xa0")}
        updated = text
        for variant in variants:
            if variant in updated:
                updated = updated.replace(variant, new)
        if updated == text:
            plain = text.replace("\xa0", " ")
            if old.replace("\xa0", " ") in plain:
                updated = plain.replace(old.replace("\xa0", " "), new)
        if updated != text:
            if para.runs:
                para.runs[0].text = updated
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.add_run(updated)
            count += 1
    return count


def replace_everywhere(doc: Document, old: str, new: str) -> int:
    return replace_all_paragraphs(doc, old, new) + replace_all_cells(doc, old, new)


def set_row_values(row, values: list[str]) -> None:
    for idx, value in enumerate(values):
        if idx < len(row.cells):
            set_cell(row.cells[idx], value)


def insert_rows_after_id(table, after_id: str, rows_values: list[list[str]]) -> int:
    inserted = 0
    matches = [i for i, row in enumerate(table.rows) if cell_text(row.cells[0]) == after_id]
    for idx in reversed(matches):
        cursor = table.rows[idx]._tr
        for values in rows_values:
            new_tr = deepcopy(table.rows[idx]._tr)
            cursor.addnext(new_tr)
            cursor = new_tr
            inserted += 1
        # Fill from the newly inserted XML rows immediately after original idx.
        # python-docx refreshes rows from XML, so original idx+1 .. idx+len are new.
    # Second pass: fill values for each match from the end after all inserts.
    # After inserting len(rows_values) after each match from the end, fill by scanning.
    return _fill_inserted(table, after_id, rows_values, inserted)


def _fill_inserted(table, after_id: str, rows_values: list[list[str]], expected: int) -> int:
    filled = 0
    i = 0
    while i < len(table.rows):
        if cell_text(table.rows[i].cells[0]) == after_id:
            for offset, values in enumerate(rows_values, start=1):
                target = i + offset
                if target < len(table.rows):
                    set_row_values(table.rows[target], values)
                    filled += 1
            i += 1 + len(rows_values)
            continue
        i += 1
    return filled


def insert_paragraph_after(paragraph, text: str, style_name: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source: {SRC}")
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    # --- Paragraph / cell text aligned with SRS 2.0.1 ---
    replacements = [
        (
            "Thanh toán, nhận Ticket QR và xem lịch sử.",
            "Thanh toán trả trước qua cổng, nhận Ticket QR và xem lịch sử.",
        ),
        (
            "Hủy Ticket/Booking và tạo Refund theo policy.",
            "Hủy Ticket/Booking; Refund nền tảng chỉ khi đã thu trả trước.",
        ),
        (
            "Operator quản lý Organization, Bus, Seat, Driver, Route, Stop và Trip.",
            "Operator quản lý Organization (kể cả bật/tắt trả sau), Bus, Seat, Driver, Route, Stop và Trip.",
        ),
        (
            "Hệ thống consume hold và tạo một Booking PENDING_PAYMENT.",
            "Hệ thống consume hold: PREPAID tạo Booking PENDING_PAYMENT; PAY_LATER (nhà xe bật) tạo Booking CONFIRMED và phát Ticket ngay, không tạo Payment cổng.",
        ),
        (
            "Hệ thống trả Booking summary, total chính thức và thời hạn thanh toán.",
            "Hệ thống trả Booking summary, total chính thức, paymentChannel và thời hạn thanh toán nếu PREPAID.",
        ),
        (
            "Customer mở Booking còn thời hạn thanh toán.",
            "Customer mở Booking PREPAID còn thời hạn thanh toán. Booking PAY_LATER/CONFIRMED không tạo Payment cổng.",
        ),
        (
            "Hệ thống tính phí, số tiền hoàn, phương thức hoàn và thời gian dự kiến.",
            "Hệ thống tính phí, số tiền hoàn (0 nếu PAY_LATER), phương thức hoàn và thời gian dự kiến.",
        ),
        (
            "Nếu refund amount lớn hơn 0, hệ thống tạo Refund REQUESTED.",
            "Nếu PREPAID đã thu và refund amount lớn hơn 0, hệ thống tạo Refund REQUESTED. PAY_LATER không hoàn qua cổng.",
        ),
        (
            "Booking PAID có đúng một Ticket cho mỗi item.",
            "Booking PAID (PREPAID) hoặc CONFIRMED (PAY_LATER) có đúng một Ticket cho mỗi item.",
        ),
        (
            "Giúp Customer tìm được Trip phù hợp, giữ ghế, tạo Booking, thanh toán và nhận Ticket mà không xảy ra bán trùng ghế.",
            "Giúp Customer tìm được Trip phù hợp, giữ ghế, tạo Booking, thanh toán trả trước hoặc trả sau khi nhà xe bật, và nhận Ticket mà không xảy ra bán trùng ghế.",
        ),
        (
            "Thu hồi quyền sử dụng Ticket đủ điều kiện và hoàn đúng số tiền cho Customer theo policy đã được chốt khi đặt vé.",
            "Thu hồi quyền sử dụng Ticket đủ điều kiện; hoàn cổng chỉ cho khoản PREPAID đã thu theo policy đã chốt khi đặt vé.",
        ),
    ]
    for old, new in replacements:
        replace_everywhere(doc, old, new)

    # UC-BOOK-01 exceptions
    replace_everywhere(
        doc,
        "Số Passenger khác số ghế: trả validation error; chưa consume hold.",
        "Số Passenger khác số ghế: trả validation error; chưa consume hold. PAY_LATER khi nhà xe tắt: PAY_LATER_NOT_ALLOWED, không consume hold.",
    )

    # UC-OPS-01 flow: allowPayLater
    replace_everywhere(
        doc,
        "Actor cập nhật tên hiển thị, liên hệ, mô tả hoặc thuộc tính được phép.",
        "Actor cập nhật tên hiển thị, liên hệ, mô tả, cho phép trả sau hoặc thuộc tính được phép. commissionRate thuộc nền tảng, PATCH tenant bỏ qua.",
    )

    # UC-TRIP-01 refund prepaid only
    replace_everywhere(
        doc,
        "Hệ thống tạo Refund cho khoản đủ điều kiện theo policy nhà xe.",
        "Hệ thống tạo Refund cổng cho khoản PREPAID đã thu; PAY_LATER chỉ hủy vé/chỗ.",
    )

    # UC-REPORT
    replace_everywhere(
        doc,
        "Hệ thống trả gross revenue, net revenue, Booking, Refund và occupancy cùng định nghĩa metric.",
        "Hệ thống trả gross/net trên tiền PREPAID đã thu, phí sàn, công nợ nhà xe, Booking, Refund và occupancy; không cộng tiền mặt PAY_LATER.",
    )

    # UC-ADMIN-02
    replace_everywhere(
        doc,
        "Hệ thống tìm và liên kết Booking, Payment, Ticket, Refund, Notification và audit liên quan.",
        "Hệ thống tìm và liên kết Booking, Payment, Ticket, Refund, settlement, Notification và audit liên quan.",
    )

    # --- Tables: BP-01 ---
    replace_everywhere(
        doc,
        "Nhập Passenger, điểm đón/trả và Promotion nếu có; hệ thống tính lại giá và tạo Booking.",
        "Nhập Passenger, điểm đón/trả; chọn PREPAID hoặc PAY_LATER (chỉ khi nhà xe bật); hệ thống tính lại giá và tạo Booking.",
    )
    replace_everywhere(
        doc,
        "Booking PENDING_PAYMENT có tổng tiền chính thức được tạo từ SeatHold.",
        "Booking PENDING_PAYMENT (trả trước) hoặc CONFIRMED (trả sau) được tạo từ SeatHold.",
    )
    replace_everywhere(
        doc,
        "Customer chọn phương thức và thực hiện thanh toán.",
        "PREPAID: Customer thanh toán qua cổng. PAY_LATER: không tạo Payment cổng; phát Ticket ngay.",
    )
    replace_everywhere(
        doc,
        "Payment được tạo và kết quả được gửi về hệ thống.",
        "PREPAID: Payment được tạo. PAY_LATER: ghế BOOKED, mỗi item một Ticket ISSUED ghi PAY_LATER.",
    )
    replace_everywhere(
        doc,
        "Xác minh provider, chữ ký, transaction ID, amount và currency.",
        "PREPAID: xác minh provider, chữ ký, transaction ID, amount và currency. PAY_LATER: bỏ qua cổng.",
    )
    replace_everywhere(
        doc,
        "Cập nhật Booking, TripSeat và phát hành Ticket theo cơ chế idempotent.",
        "PREPAID: cập nhật Booking PAID, ghế, Ticket và chốt phí sàn. PAY_LATER: vé đã phát lúc CONFIRMED.",
    )
    replace_everywhere(
        doc,
        "Booking PAID, TripSeat BOOKED và mỗi Booking Item có một Ticket ISSUED.",
        "PREPAID: Booking PAID, settlement COLLECTED. PAY_LATER: Ticket đã ISSUED, phí sàn = 0.",
    )

    replace_everywhere(
        doc,
        "Payment chưa có kết quả cuối, webhook lặp hoặc dữ liệu xác minh không hợp lệ.",
        "Nhà xe không bật trả sau nhưng client gửi PAY_LATER; hoặc Payment PREPAID chưa có kết quả cuối/webhook lặp/mismatch.",
    )
    replace_everywhere(
        doc,
        "Giữ trạng thái phù hợp để truy vấn/đối soát; không xác nhận Booking hoặc tạo Ticket lặp.",
        "PAY_LATER_NOT_ALLOWED không consume hold. PREPAID: không xác nhận Booking hoặc tạo Ticket lặp.",
    )

    # BP-02
    replace_everywhere(
        doc,
        "Kiểm tra ownership, trạng thái, giờ khởi hành và policy snapshot; tính phí và số tiền hoàn.",
        "Kiểm tra ownership, trạng thái, giờ khởi hành, paymentChannel và policy snapshot; tính phí và số tiền hoàn.",
    )
    replace_everywhere(
        doc,
        "Preview hủy vé được hiển thị mà chưa thay đổi trạng thái.",
        "Preview: PREPAID có thể có refund; PAY_LATER hoàn nền tảng = 0. Chưa đổi trạng thái.",
    )
    replace_everywhere(
        doc,
        "Tạo và xử lý Refund nếu số tiền hoàn lớn hơn 0.",
        "Chỉ tạo Refund cổng khi PREPAID đã thu và số hoàn > 0.",
    )
    replace_everywhere(
        doc,
        "Refund được theo dõi đến trạng thái cuối hoặc trạng thái cần retry.",
        "PREPAID: theo dõi Refund. PAY_LATER: hủy chỗ/vé, không hoàn cổng.",
    )

    # BR updates
    replace_everywhere(
        doc,
        "Mỗi ghế trong Booking có đúng một Passenger và sau thanh toán có đúng một Ticket.",
        "Mỗi ghế trong Booking có đúng một Passenger. Ticket phát hành khi PREPAID đã PAID hoặc ngay khi PAY_LATER chuyển CONFIRMED.",
    )
    replace_everywhere(
        doc,
        "Booking PENDING_PAYMENT quá hạn mà chưa có Payment hợp lệ chuyển EXPIRED.",
        "Booking PREPAID PENDING_PAYMENT quá hạn chưa có Payment hợp lệ chuyển EXPIRED. PAY_LATER không dùng cửa sổ thanh toán cổng.",
    )
    replace_everywhere(
        doc,
        "Booking chỉ chuyển PAID sau kết quả Payment thành công đã được xác minh.",
        "Booking chỉ chuyển PAID sau Payment cổng thành công đã xác minh. PAY_LATER không đi qua PAID.",
    )

    # UC attribute tables
    replace_everywhere(
        doc,
        "Giữ toàn bộ ghế đã chọn và tạo một Booking chờ thanh toán.",
        "Giữ toàn bộ ghế đã chọn và tạo Booking PREPAID chờ thanh toán hoặc PAY_LATER CONFIRMED kèm vé.",
    )
    replace_everywhere(
        doc,
        "Booking PENDING_PAYMENT; SeatHold/TripSeat được bảo toàn đến thời hạn thanh toán.",
        "PREPAID: Booking PENDING_PAYMENT. PAY_LATER (nhà xe bật): Booking CONFIRMED, ghế BOOKED, mỗi item một Ticket PAY_LATER, không Payment cổng.",
    )
    replace_everywhere(
        doc,
        "FR-BOOK-001..007; BR-SEAT-; BR-BOOK-; AC-SEAT-; AC-BOOK-",
        "FR-BOOK-001..007, FR-BOOK-012; BR-SEAT-*; BR-BOOK-*; AC-SEAT-*; AC-BOOK-*",
    )
    replace_everywhere(
        doc,
        "Booking PENDING_PAYMENT, còn hạn và thuộc Customer.",
        "Booking PREPAID PENDING_PAYMENT, còn hạn và thuộc Customer. PAY_LATER/CONFIRMED: từ chối create Payment cổng.",
    )
    replace_everywhere(
        doc,
        "FR-PAY-001..007; FR-TICKET-001..003; BR-PAY-; AC-PAY-",
        "FR-PAY-001..007, FR-PAY-012; FR-TICKET-001..003; BR-PAY-*; AC-PAY-*",
    )
    replace_everywhere(
        doc,
        "Thu hồi quyền sử dụng Ticket và hoàn đúng số tiền theo policy.",
        "Thu hồi quyền sử dụng Ticket; hoàn cổng chỉ khoản PREPAID đã thu.",
    )
    replace_everywhere(
        doc,
        "Ticket bị hủy; ghế được mở lại nếu còn bán; Refund được tạo khi có tiền phải hoàn.",
        "Ticket bị hủy; ghế mở lại nếu còn bán; Refund cổng chỉ khi PREPAID đã thu và số hoàn > 0.",
    )
    replace_everywhere(
        doc,
        "FR-BOOK-009; FR-PAY-008; BR-CANCEL-; AC-CANCEL-",
        "FR-BOOK-009; FR-PAY-008, FR-PAY-011; BR-CANCEL-*; AC-CANCEL-*",
    )
    replace_everywhere(
        doc,
        "FR-OPS-001; BR-TENANT-*; AUTHZ-002..003",
        "FR-OPS-001, FR-OPS-011; BR-TENANT-*; BR-TRIP-007; AUTHZ-002..003",
    )
    replace_everywhere(
        doc,
        "Dừng Trip, vô hiệu Ticket liên quan và khởi tạo Refund an toàn.",
        "Dừng Trip, vô hiệu Ticket liên quan; hoàn cổng chỉ khoản PREPAID đã thu.",
    )
    replace_everywhere(
        doc,
        "Trip CANCELLED; Ticket không còn sử dụng; Refund được tạo theo policy.",
        "Trip CANCELLED; Ticket không còn sử dụng; Refund cổng chỉ PREPAID đã thu; PAY_LATER hủy vé không hoàn cổng.",
    )

    # FR cell updates (all copies)
    replace_everywhere(
        doc,
        "Customer có thể tạo đúng một Booking từ SeatHold còn hiệu lực; thao tác hỗ trợ idempotency.",
        "Customer có thể tạo đúng một Booking từ SeatHold còn hiệu lực, chọn PREPAID hoặc PAY_LATER khi nhà xe cho phép; thao tác hỗ trợ idempotency.",
    )
    replace_everywhere(
        doc,
        "SeatHold/Booking chưa thanh toán hết hạn được chuyển trạng thái và giải phóng ghế tự động.",
        "SeatHold/Booking PREPAID chưa thanh toán hết hạn được chuyển trạng thái và giải phóng ghế tự động.",
    )
    replace_everywhere(
        doc,
        "Customer có thể hủy toàn Booking hoặc Ticket đủ điều kiện; hệ thống hiển thị phí và số tiền hoàn trước khi xác nhận.",
        "Customer có thể hủy toàn Booking hoặc Ticket đủ điều kiện; preview phải ghi rõ số hoàn (0 nếu PAY_LATER).",
    )
    replace_everywhere(
        doc,
        "Customer có thể tạo Payment intent cho Booking PENDING_PAYMENT còn hiệu lực.",
        "Customer có thể tạo Payment intent cho Booking PREPAID PENDING_PAYMENT còn hiệu lực.",
    )
    replace_everywhere(
        doc,
        "Hệ thống tạo và theo dõi Refund; refund request lặp phải idempotent.",
        "Hệ thống tạo và theo dõi Refund cổng; refund request lặp phải idempotent; từ chối Refund nếu không có Payment SUCCEEDED.",
    )
    replace_everywhere(
        doc,
        "Admin/Operator Finance có thể tra cứu Payment/Refund theo phạm vi quyền và mã giao dịch.",
        "Admin/Operator Finance có thể tra cứu Payment/Refund/settlement theo phạm vi quyền và mã giao dịch.",
    )
    replace_everywhere(
        doc,
        "Mỗi Passenger/TripSeat của Booking PAID có đúng một Ticket được phát hành.",
        "Mỗi Booking Item có đúng một Ticket khi Booking PAID (PREPAID) hoặc CONFIRMED (PAY_LATER). Vé hiển thị kênh thanh toán.",
    )
    replace_everywhere(
        doc,
        "Ticket hiển thị mã vé, Passenger, nhà xe, Trip, điểm đón/trả, ghế, giá snapshot, trạng thái và QR.",
        "Ticket hiển thị mã vé, Passenger, nhà xe, Trip, điểm đón/trả, ghế, giá snapshot, paymentChannel, trạng thái và QR.",
    )
    replace_everywhere(
        doc,
        "Operator Staff có permission phù hợp có thể cập nhật thông tin Organization của mình.",
        "Operator Staff có permission phù hợp có thể cập nhật thông tin Organization của mình, gồm bật/tắt cho phép trả sau.",
    )
    replace_everywhere(
        doc,
        "Hủy Trip có Ticket đã bán phải khởi tạo xử lý hủy vé, Refund và Notification.",
        "Hủy Trip có Ticket đã bán phải vô hiệu vé, thông báo khách, và chỉ tạo Refund cổng cho khoản PREPAID đã thu.",
    )
    replace_everywhere(
        doc,
        "Hệ thống tạo Notification cho Booking paid, Ticket issued/changed, Payment failed, Trip changed/cancelled, Booking cancelled và Refund completed.",
        "Hệ thống tạo Notification cho Booking paid/CONFIRMED, Ticket issued/changed, Payment failed, Trip changed/cancelled, Booking cancelled và Refund completed.",
    )
    replace_everywhere(
        doc,
        "Admin xem gross/net revenue, Booking, Refund và occupancy theo khoảng thời gian.",
        "Admin xem gross/net revenue, phí sàn, Booking, Refund và occupancy theo khoảng thời gian. Net nhà xe trên tiền PREPAID đã thu không gồm tiền mặt trả sau.",
    )

    # Data chapter entities
    replace_everywhere(
        doc,
        "ID/code, Customer ID, Trip ID, status, subtotal, discount, fee, total, currency, expiry",
        "ID/code, Customer ID, Trip ID, status, paymentChannel, subtotal, discount, fee, total, currency, expiry",
    )
    replace_everywhere(
        doc,
        "Code duy nhất; tiền chính xác; giữ policy snapshot.",
        "Code duy nhất; paymentChannel PREPAID/PAY_LATER; PAY_LATER chỉ khi org cho phép.",
    )
    replace_everywhere(
        doc,
        "Một Ticket có hiệu lực cho mỗi Booking Item/TripSeat.",
        "Một Ticket/item; paymentChannel; PAY_LATER không chuyển REFUNDED.",
    )

    # Ticket state
    replace_everywhere(
        doc,
        "Khoản hoàn liên quan thành công.",
        "Khoản hoàn cổng PREPAID thành công. Ticket PAY_LATER hủy xong dừng ở CANCELLED.",
    )

    # Booking PAID cancel note
    replace_everywhere(
        doc,
        "Có khoản đã thanh toán cần hoàn.",
        "PREPAID có khoản đã thu cần hoàn. PAY_LATER hủy/no-show dừng ở CANCELLED, không REFUND_PENDING.",
    )

    # AC-TICKET / AC-TRIP
    replace_everywhere(
        doc,
        "Booking PAID có nhiều Booking Item",
        "Booking PAID hoặc CONFIRMED có nhiều Booking Item",
    )
    replace_everywhere(
        doc,
        "Mỗi item có đúng một Ticket, không trùng.",
        "Mỗi item có đúng một Ticket, không trùng; Ticket có paymentChannel.",
    )
    replace_everywhere(
        doc,
        "Có đủ Trip, Passenger, ghế, điểm đón/trả, giá, trạng thái, QR và public code an toàn.",
        "Có đủ Trip, Passenger, ghế, điểm đón/trả, giá, kênh thanh toán, trạng thái, QR và public code an toàn.",
    )
    replace_everywhere(
        doc,
        "Trip CANCELLED, Ticket bị vô hiệu, Refund được yêu cầu và Customer được thông báo eventual.",
        "Trip CANCELLED, Ticket bị vô hiệu; Refund cổng chỉ cho PREPAID đã thu; PAY_LATER hủy vé không hoàn cổng; Customer được thông báo eventual.",
    )

    # MUST bullets: trả sau + phí sàn sit with booking/payment scope
    for para in doc.paragraphs:
        if para.text.strip().startswith("Hủy Ticket/Booking"):
            nxt = insert_paragraph_after(
                para,
                "Trả sau khi nhà xe bật: in vé ghi rõ kênh trả; sàn không theo dõi tiền mặt trên xe; khách không tới do nhà xe chịu.",
                para.style.name if para.style else "Normal",
            )
            insert_paragraph_after(
                nxt,
                "Phí sàn (mặc định 10%) trừ trên tiền đã vào cổng; phần còn lại ghi công nợ trả nhà xe.",
                para.style.name if para.style else "Normal",
            )
            break

    # Product principles
    for para in doc.paragraphs:
        if para.text.strip().startswith("Payment redirect không phải bằng chứng"):
            nxt = insert_paragraph_after(
                para,
                "paymentChannel là kênh hợp đồng lúc đặt (PREPAID / PAY_LATER), không phải thời điểm đo tiền mặt trên xe.",
                para.style.name if para.style else "Normal",
            )
            insert_paragraph_after(
                nxt,
                "Nền tảng chỉ hoàn tiền khi đã thu qua cổng; phí sàn chỉ tính trên khoản PREPAID đã SUCCEEDED.",
                para.style.name if para.style else "Normal",
            )
            break

    # Version history after 1.5
    for para in doc.paragraphs:
        if para.text.strip() == "1.5 Quy ước yêu cầu":
            h = insert_paragraph_after(para, "1.6 Lịch sử phiên bản nội dung", "Heading 2")
            p1 = insert_paragraph_after(
                h,
                "Phiên bản nội dung SRS và phiên bản file Word là hai trục độc lập. Markdown docs/srs/v2 là nguồn quyết định hành vi.",
                "Normal",
            )
            p2 = insert_paragraph_after(
                p1,
                "2.0.0 (09/09/2026): baseline MVP; đặt vé thanh toán trước qua cổng.",
                "List Paragraph",
            )
            p3 = insert_paragraph_after(
                p2,
                "2.0.1 (17/09/2026): bổ sung MUST trả sau (nhà xe bật, in vé, sàn không đo tiền mặt, no-show do nhà xe chịu) và phí sàn mặc định 10% trên tiền PREPAID đã thu qua cổng. Word v002 đóng gói baseline này.",
                "List Paragraph",
            )
            insert_paragraph_after(
                p3,
                "Bản Word này là v002, phát hành 17/09/2026, supersede v001 về nội dung nghiệp vụ; v001 được giữ nguyên trong releases/.",
                "Normal",
            )
            break

    # New FR / BR / AC / entity / error rows
    insert_rows_after_id(
        doc.tables[25],
        "BR-BOOK-010",
        [
            [
                "BR-BOOK-011",
                "paymentChannel chỉ PREPAID hoặc PAY_LATER. PAY_LATER chỉ khi Organization đang bật cho phép trả sau. Đây là nhãn hợp đồng lúc đặt, không phải timestamp thu tiền mặt.",
            ],
            [
                "BR-BOOK-012",
                "PAY_LATER: consume hold, ghế BOOKED, phát Ticket ngay, Booking CONFIRMED; không tạo Payment cổng.",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[26],
        "BR-PAY-010",
        [
            [
                "BR-PAY-011",
                "Nền tảng chỉ tạo Refund cổng cho Booking/Ticket PREPAID đã thu. PAY_LATER không hoàn qua nền tảng.",
            ],
            [
                "BR-PAY-012",
                "Khi Payment PREPAID SUCCEEDED, chốt phí sàn (snapshot tỷ lệ Organization, mặc định 10%) và công nợ nhà xe = gross − commission.",
            ],
            [
                "BR-PAY-013",
                "Hoàn PREPAID đảo commission và operator payable tương ứng, không vượt số đã thu.",
            ],
            [
                "BR-PAY-014",
                "PAY_LATER no-show: không Refund, phí sàn = 0; nhà xe chịu chỗ trống.",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[64],
        "FR-BOOK-011",
        [
            [
                "FR-BOOK-012",
                "MUST",
                "Khi PAY_LATER hợp lệ, hệ thống phát hành Ticket ngay, ghế BOOKED, Booking CONFIRMED; không tạo Payment cổng.",
                "UC-BOOK-01",
            ]
        ],
    )
    insert_rows_after_id(
        doc.tables[65],
        "FR-PAY-010",
        [
            [
                "FR-PAY-011",
                "MUST",
                "Nền tảng chỉ hoàn tiền các khoản đã thu PREPAID; PAY_LATER không hoàn qua cổng.",
                "UC-CANCEL-01",
            ],
            [
                "FR-PAY-012",
                "MUST",
                "Khi Payment PREPAID thành công, hệ thống chốt phí sàn (snapshot tỷ lệ Organization, mặc định 10%) và công nợ nhà xe = gross − commission.",
                "UC-PAY-01, UC-REPORT-01",
            ],
            [
                "FR-PAY-013",
                "MUST",
                "Operator Finance/Admin xem được settlement/payout; hoàn PREPAID đảo commission và payable.",
                "UC-ADMIN-02, UC-REPORT-01",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[67],
        "FR-OPS-010",
        [
            [
                "FR-OPS-011",
                "MUST",
                "Tỷ lệ phí sàn thuộc nền tảng; Operator không đổi commission qua API tenant.",
                "UC-OPS-01, UC-ADMIN-01",
            ]
        ],
    )
    confirm_rows = [
        ["—", "CONFIRMED", "Tạo Booking PAY_LATER thành công: ghế BOOKED và Ticket đã phát."],
        ["CONFIRMED", "CANCELLED", "Hủy/no-show PAY_LATER; không Refund cổng."],
        ["CONFIRMED", "COMPLETED", "Trip hoàn thành với Booking trả sau."],
    ]
    table72 = doc.tables[72]
    header_idxs = [
        i
        for i, row in enumerate(table72.rows)
        if cell_text(row.cells[1]) == "Đến" and cell_text(row.cells[0]) in ("Từ", "")
    ]
    for idx in reversed(header_idxs):
        cursor = table72.rows[idx]._tr
        for _values in confirm_rows:
            new_tr = deepcopy(table72.rows[idx]._tr)
            cursor.addnext(new_tr)
            cursor = new_tr
    i = 0
    while i < len(table72.rows):
        first = cell_text(table72.rows[i].cells[0])
        second = cell_text(table72.rows[i].cells[1])
        if second == "Đến" and first in ("Từ", ""):
            for offset, values in enumerate(confirm_rows, start=1):
                if i + offset >= len(table72.rows):
                    break
                nxt_first = cell_text(table72.rows[i + offset].cells[0])
                nxt_second = cell_text(table72.rows[i + offset].cells[1])
                if nxt_second == "Đến" or nxt_first in ("Từ", ""):
                    set_row_values(table72.rows[i + offset], values)
            i += 1 + len(confirm_rows)
            continue
        i += 1

    insert_rows_after_id(
        doc.tables[82],
        "ReconciliationCase",
        [
            [
                "BookingSettlement",
                "Booking ID, Organization ID, paymentChannel, gross, commissionRate/amount, operatorNet, collectionStatus",
                "PREPAID đã thu: commission + operatorNet = gross. PAY_LATER: commission = 0, không Payment cổng.",
            ],
            [
                "LedgerEntry",
                "Settlement ID, entryType, amount, correlationId",
                "Append-only; capture/commission/payable và đảo khi hoàn PREPAID.",
            ],
            [
                "OperatorPayout",
                "Organization ID, period, payableAmount, status",
                "Ghi kỳ chuyển net PREPAID cho nhà xe.",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[97],
        "AC-BOOK-005",
        [
            [
                "AC-BOOK-006",
                "Organization tắt trả sau",
                "Customer tạo Booking PAY_LATER",
                "Từ chối PAY_LATER_NOT_ALLOWED; không consume hold.",
            ],
            [
                "AC-BOOK-007",
                "Organization bật trả sau; hold còn hạn",
                "Customer tạo Booking PAY_LATER",
                "Booking CONFIRMED, mỗi item một Ticket PAY_LATER, ghế BOOKED, không có Payment cổng.",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[98],
        "AC-PAY-006",
        [
            [
                "AC-PAY-007",
                "Booking PAY_LATER CONFIRMED",
                "Customer yêu cầu Refund nền tảng hoặc create Payment",
                "Từ chối; không tạo Payment/Refund cổng.",
            ],
            [
                "AC-PAY-008",
                "Payment PREPAID SUCCEEDED, org commission 10%, gross 100000",
                "Chốt settlement",
                "commission 10000, operator_net 90000, collection COLLECTED.",
            ],
            [
                "AC-PAY-009",
                "Refund PREPAID thành công một phần/toàn bộ",
                "Cập nhật settlement",
                "Đảo commission/payable tương ứng; không vượt số đã thu.",
            ],
        ],
    )
    insert_rows_after_id(
        doc.tables[103],
        "RATE_LIMITED",
        [
            [
                "PAY_LATER_NOT_ALLOWED",
                "409",
                "Nhà xe không bật trả sau",
                "Không tạo Booking PAY_LATER.",
            ]
        ],
    )

    # BP-02 extra exception
    insert_rows_after_id(
        doc.tables[12],
        "5",
        [
            [
                "6",
                "Ticket PAY_LATER được hủy hoặc no-show.",
                "Không tạo Payment/Refund cổng; nhà xe chịu ghế trống.",
            ]
        ],
    )
    # The STT "5" matches multiple tables' rows; table 12 is targeted. Also "5" as first cell exists in T10.
    # insert_rows_after_id on T12 only — good.

    insert_rows_after_id(
        doc.tables[10],
        "6",
        [
            [
                "7",
                "PAY_LATER: khách không lên xe.",
                "Hủy quyền vé/chỗ theo cutoff; không Refund nền tảng; nhà xe chịu ghế trống.",
            ]
        ],
    )

    # Cover page date
    replace_everywhere(doc, "Đà Nẵng, 06/2026", "Đà Nẵng, 09/2026")

    now = datetime(2026, 9, 17, tzinfo=timezone.utc)
    doc.core_properties.modified = now
    doc.core_properties.title = "PBL6 SRS v002 — baseline nội dung 2.0.1"
    doc.core_properties.comments = (
        "Word v002 packages SRS 2.0.1: PAY_LATER and 10% prepaid platform commission. "
        "Does not overwrite v001."
    )
    doc.core_properties.revision = 2
    doc.save(str(DST))

    digest = hashlib.sha256(DST.read_bytes()).hexdigest().upper()
    size = DST.stat().st_size
    print(f"Wrote {DST}")
    print(f"Size {size}")
    print(f"SHA-256 {digest}")


if __name__ == "__main__":
    main()
