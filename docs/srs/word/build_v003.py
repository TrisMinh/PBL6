"""Build PBL6-SRS-v003.docx from v002 without overwriting previous releases."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "docs/srs/word/releases/PBL6-SRS-v002.docx"
DST = ROOT / "docs/srs/word/releases/PBL6-SRS-v003.docx"


def plain(value: str) -> str:
    return value.replace("\xa0", " ").strip()


def set_paragraph(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    set_paragraph(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        parent = extra._element.getparent()
        if parent is not None:
            parent.remove(extra._element)


def set_row(row, values: list[str]) -> None:
    for index, value in enumerate(values):
        if index < len(row.cells):
            set_cell(row.cells[index], value)


def paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    created = Paragraph(element, paragraph._parent)
    if style:
        created.style = style
    created.add_run(text)
    return created


def insert_sequence_after(paragraph: Paragraph, items: list[tuple[str, str]]) -> Paragraph:
    cursor = paragraph
    for text, style in items:
        cursor = paragraph_after(cursor, text, style)
    return cursor


def find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if plain(paragraph.text) == exact_text:
            return paragraph
    raise RuntimeError(f"Paragraph not found: {exact_text}")


def replace_paragraph(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if plain(paragraph.text) == old:
            set_paragraph(paragraph, new)
            count += 1
    return count


def replace_cells(doc: Document, old: str, new: str) -> int:
    count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if plain(cell.text) == old:
                    set_cell(cell, new)
                    count += 1
    return count


def insert_rows_after_id(table, row_id: str, rows_to_add: list[list[str]]) -> int:
    matches = [row for row in table.rows if plain(row.cells[0].text) == row_id]
    inserted = 0
    for source_row in reversed(matches):
        cursor = source_row._tr
        for values in rows_to_add:
            new_xml_row = deepcopy(source_row._tr)
            cursor.addnext(new_xml_row)
            cursor = new_xml_row
            inserted += 1

    # Fill each inserted run immediately following every matching source row.
    index = 0
    while index < len(table.rows):
        if plain(table.rows[index].cells[0].text) == row_id:
            for offset, values in enumerate(rows_to_add, start=1):
                set_row(table.rows[index + offset], values)
            index += len(rows_to_add) + 1
        else:
            index += 1
    return inserted


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source: {SRC}")
    if DST.exists():
        raise SystemExit(f"Refusing to overwrite existing release: {DST}")

    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    # Version history: retain prior baselines and append the new content baseline.
    version_anchor = find_paragraph(
        doc,
        "2.0.1 (17/09/2026): bổ sung MUST trả sau (nhà xe bật, in vé, sàn không đo tiền mặt, "
        "no-show do nhà xe chịu) và phí sàn mặc định 10% trên tiền PREPAID đã thu qua cổng. "
        "Word v002 đóng gói baseline này.",
    )
    insert_sequence_after(
        version_anchor,
        [
            (
                "2.1.0 (21/09/2026): bổ sung FR-REPORT-004 (SHOULD) cho Power BI Embedded đa chiều, "
                "tenant/RLS, freshness, UI/data/integration và AC-REPORT-004..006. 61 FR MUST của MVP không đổi.",
                "List Paragraph",
            ),
            (
                "Bản Word này là v003, phát hành 21/09/2026, đóng gói baseline nội dung 2.1.0; "
                "v001 và v002 được giữ nguyên trong releases/.",
                "Normal",
            ),
        ],
    )
    replace_paragraph(
        doc,
        "Bản Word này là v002, phát hành 17/09/2026, supersede v001 về nội dung nghiệp vụ; v001 được giữ nguyên trong releases/.",
        "Bản Word v002 đóng gói baseline 2.0.1 và đã được v003 thay thế; v001/v002 được giữ nguyên trong releases/.",
    )

    # Product scope and use-case behavior.
    insert_sequence_after(
        find_paragraph(doc, "Export báo cáo CSV."),
        [
            (
                "Dashboard đa chiều nhúng bằng Power BI, giới hạn theo permission/tenant và hiển thị độ mới dữ liệu.",
                "- lv1",
            )
        ],
    )
    replace_cells(
        doc,
        "Payment Gateway: tiếp nhận yêu cầu thanh toán, hoàn tiền và gửi webhook có chữ ký. "
        "Notification Provider: gửi email, SMS hoặc thông báo đẩy; không có quyền quyết định trạng thái nghiệp vụ của hệ thống.",
        "Payment Gateway: tiếp nhận thanh toán/hoàn tiền và gửi webhook có chữ ký. "
        "Notification Provider: gửi email, SMS hoặc push. Power BI Service: nhận dữ liệu phân tích chỉ đọc từ "
        "Reporting read model và cung cấp báo cáo nhúng; các hệ thống ngoài không quyết định trạng thái nghiệp vụ.",
    )
    insert_sequence_after(
        find_paragraph(doc, "Actor có thể drill-down/tra cứu giao dịch khi có permission."),
        [
            ("Luồng phân tích Power BI (SHOULD)", "Normal"),
            ("Actor mở dashboard Power BI từ Back-office khi feature đã được kích hoạt.", "- lv1"),
            (
                "Hệ thống xác thực actor, kiểm tra permission và suy ra tenant từ identity context; "
                "client không tự chọn tenant/report ID ngoài allowlist.",
                "- lv1",
            ),
            (
                "Hệ thống cấp cấu hình nhúng ngắn hạn, chỉ đọc và áp RLS; platform Admin chỉ xem "
                "toàn nền tảng khi có permission tương ứng.",
                "- lv1",
            ),
            (
                "Power BI đọc semantic model từ Reporting read model, không truy vấn trực tiếp database giao dịch.",
                "- lv1",
            ),
            (
                "Dashboard hiển thị metric definition, timezone, phạm vi, source data as-of và semantic-model refresh time.",
                "- lv1",
            ),
        ],
    )
    insert_sequence_after(
        find_paragraph(doc, "Operator cố xem tenant khác: từ chối."),
        [
            (
                "Power BI/RLS/effective identity thiếu hoặc không hợp lệ: fail closed, không fallback sang quyền xem toàn bộ.",
                "- lv1",
            ),
            (
                "Power BI/gateway/refresh tạm lỗi: hiển thị lỗi/fallback báo cáo cơ bản; không ảnh hưởng Booking/Payment đã commit.",
                "- lv1",
            ),
        ],
    )

    # Back-office UI and external integration.
    insert_sequence_after(
        find_paragraph(doc, "Báo cáo và Export Job."),
        [("Power BI Analytics khi feature SHOULD được kích hoạt; report theo permission và tenant/RLS.", "- lv1")],
    )
    integration_heading = find_paragraph(doc, "7.10. Tương thích hợp đồng")
    set_paragraph(integration_heading, "7.10. Power BI Service (SHOULD)")
    insert_sequence_after(
        integration_heading,
        [
            (
                "Back-office chỉ nhận embed URL/token ngắn hạn sau khi backend kiểm tra access token, permission và tenant scope.",
                "- lv1",
            ),
            (
                "Workspace/report/semantic-model ID lấy từ allowlist server; frontend không yêu cầu tùy ý ID hoặc role.",
                "- lv1",
            ),
            (
                "Operator dùng RLS theo tenant; platform Admin chỉ dùng scope toàn nền tảng khi có permission được phê duyệt.",
                "- lv1",
            ),
            (
                "Power BI chỉ đọc Reporting read model/view được duyệt, không đọc trực tiếp hoặc ghi ngược database giao dịch.",
                "- lv1",
            ),
            (
                "Data-source/service-principal credential và embed token không được commit, log hoặc persist phía client.",
                "- lv1",
            ),
            (
                "Report filter, hidden page và URL parameter không phải security boundary; lỗi embed/refresh không ảnh hưởng giao dịch.",
                "- lv1",
            ),
            ("7.11. Tương thích hợp đồng", "Heading 2"),
        ],
    )

    # Keep the visible TOC useful even before Word refreshes fields, and request
    # an automatic field/TOC refresh when the document is opened.
    for paragraph in doc.paragraphs:
        if plain(paragraph.text).startswith("7.10. Tương thích hợp đồng") and paragraph.style.name.lower().startswith("toc"):
            set_paragraph(paragraph, "7.10. Power BI Service (SHOULD)\t130")
            paragraph_after(paragraph, "7.11. Tương thích hợp đồng\t130", paragraph.style.name)
            break
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    insert_sequence_after(
        find_paragraph(doc, "Export, Report và Audit cũng phải áp tenant scope."),
        [
            (
                "Power BI semantic model/report phải áp tenant scope bằng RLS hoặc isolation tương đương; "
                "filter/slicer giao diện không phải ranh giới bảo mật.",
                "- lv1",
            )
        ],
    )

    # Keep repeated styled table sections in sync with the Markdown baseline.
    replace_cells(doc, "FR-REPORT-001..003; GOAL-007", "FR-REPORT-001..004; GOAL-007")
    insert_rows_after_id(
        doc.tables[77],
        "UI-010",
        [
            [
                "UI-011",
                "Dashboard Power BI nhúng có loading/error/retry/token-expired state, hiển thị timezone và độ mới dữ liệu; "
                "hidden page/filter không phải authorization.",
            ]
        ],
    )
    insert_rows_after_id(
        doc.tables[69],
        "FR-REPORT-003",
        [
            [
                "FR-REPORT-004",
                "SHOULD",
                "Admin/Operator Finance có quyền xem dashboard đa chiều Power BI nhúng trong Back-office, "
                "chỉ từ Reporting read model; áp permission và tenant/RLS, không lộ PII/credential, "
                "hiển thị metric, timezone, phạm vi và độ mới dữ liệu.",
                "UC-REPORT-01",
            ]
        ],
    )
    insert_rows_after_id(
        doc.tables[83],
        "ExportJob",
        [
            [
                "BI Semantic Model",
                "Fact/dimension/measure, tenant key, sourceDataAsOf, modelRefreshedAt",
                "Bản sao phân tích chỉ đọc từ Reporting read model; không chứa PII không cần thiết, credential hoặc dữ liệu ngoài scope.",
            ]
        ],
    )
    insert_rows_after_id(
        doc.tables[101],
        "AC-REPORT-003",
        [
            [
                "AC-REPORT-004",
                "Power BI bật; Operator Finance tenant A có permission",
                "Mở dashboard Power BI",
                "Chỉ dữ liệu tenant A; hiển thị metric, timezone, source data as-of và semantic-model refresh time.",
            ],
            [
                "AC-REPORT-005",
                "Actor không có permission, tenant/effective identity sai hoặc token hết hạn",
                "Yêu cầu/mở dashboard Power BI",
                "Fail closed; không cấp quyền toàn nền tảng, không lộ report/credential/token/PII; audit an toàn.",
            ],
            [
                "AC-REPORT-006",
                "Power BI API, gateway hoặc refresh tạm lỗi",
                "Actor dùng hệ thống",
                "Dashboard báo lỗi/freshness và cho retry/fallback; Booking/Payment đã commit không bị ảnh hưởng.",
            ],
        ],
    )

    now = datetime(2026, 9, 21, tzinfo=timezone.utc)
    doc.core_properties.modified = now
    doc.core_properties.title = "PBL6 SRS v003 — baseline nội dung 2.1.0"
    doc.core_properties.comments = (
        "Word v003 packages SRS 2.1.0: Power BI Embedded analytics is a SHOULD capability after MVP; "
        "61 MUST requirements remain unchanged. Does not overwrite v002."
    )
    doc.core_properties.revision = 3
    doc.save(str(DST))

    digest = hashlib.sha256(DST.read_bytes()).hexdigest().upper()
    print(f"Wrote {DST}")
    print(f"Size {DST.stat().st_size}")
    print(f"SHA-256 {digest}")


if __name__ == "__main__":
    main()
