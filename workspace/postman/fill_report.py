from pathlib import Path
from shutil import copyfile
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment
src=Path('Test_Report_API_Postman.xlsx')
out=Path('Test_Report_API_Postman_Result.xlsx')
copyfile(src,out)
wb=openpyxl.load_workbook(out)
executed={
 'AUTH-08':('Passed','Đã chạy Postman CLI 1.80: POST /api/v1/auth/login bằng seed customer trả HTTP 200 và có accessToken + refreshToken.'),
 'AUTH-09':('Passed','Đã chạy Postman CLI 1.80: đăng nhập bằng số điện thoại seed +840000000001 trả HTTP 200 và có accessToken.'),
 'AUTH-10':('Passed','Đã chạy Postman CLI 1.80: POST /api/v1/auth/login với mật khẩu sai trả HTTP 401; assertion không cấp phiên được xác nhận.'),
 'AUTH-13':('Passed','Đã chạy Postman CLI 1.80: refresh token hợp lệ trả HTTP 200 và token mới; dùng lại token cũ trả HTTP 401.'),
 'PRF-01':('Passed','Đã chạy Postman CLI 1.80 với access token seed customer: GET /api/v1/users/me trả HTTP 200 và đúng email của tài khoản.'),
 'CMN-03':('Passed','Đã chạy Postman CLI 1.80: request có X-Correlation-Id trả HTTP 200 và response có X-Correlation-Id.'),
 'SRCH-01':('Passed','Đã chạy Postman CLI 1.80 theo OpenAPI thực tế: GET /api/v1/trips với seed route/date trả HTTP 200, có chuyến sellable. Workbook ghi /trips/search; path được đối chiếu sang /trips.'),
 'SRCH-03':('Passed','Đã chạy Postman CLI 1.80 theo OpenAPI thực tế: tìm route không tồn tại trả HTTP 200 với items rỗng.'),
 'SRCH-08':('Failed','Đã chạy Postman CLI 1.80: size=100000 trả HTTP 200, không trả 400 như ca kiểm thử yêu cầu; OpenAPI giới hạn size tối đa 100. Assertion thất bại.'),
 'SRCH-09':('Passed','Đã chạy Postman CLI 1.80: GET /api/v1/trips/{tripId} với seed trip trả HTTP 200 và đúng ID.'),
 'SRCH-10':('Passed','Đã chạy Postman CLI 1.80: GET /api/v1/trips/{id không tồn tại} trả HTTP 404.'),
 'SRCH-11':('Passed','Đã chạy Postman CLI 1.80 với access token seed customer: GET /api/v1/trips/{tripId}/seats trả HTTP 200 và có availabilityAsOf.'),
 'BOOK-22':('Passed','Đã chạy Postman CLI 1.80 với access token seed customer: GET /api/v1/bookings trả HTTP 200 với items và totalElements.'),
 'BOOK-06':('Passed','Đã chạy Postman CLI 1.80: POST /api/v1/trips/{tripId}/seat-holds không có token trả HTTP 401.'),
}
count=0
for sh in wb.worksheets[3:]:
    for r in range(11,sh.max_row+1):
        tcid=sh.cell(r,1).value
        if not isinstance(tcid,str) or '-' not in tcid: continue
        count+=1
        status,note=executed.get(tcid,('Pending','Chưa chạy ca này trong Postman CLI ở Round 1; không suy diễn kết quả từ dotnet test hoặc đọc mã nguồn.'))
        sh.cell(r,6).value=status
        sh.cell(r,7).value=datetime(2026,10,28) if tcid in executed else None
        sh.cell(r,8).value='Postman CLI 1.80' if tcid in executed else None
        sh.cell(r,15).value=note
        sh.cell(r,15).alignment=Alignment(wrap_text=True,vertical='top')
# Correct report metadata and document the discrepancy in case count.
wb['Cover']['F6']='1.2'
wb['Cover']['A12']=datetime(2026,10,28)
wb['Cover']['B12']='1.2'
wb['Cover']['C12']='Kết quả Round 1'
wb['Cover']['D12']='M'
wb['Cover']['E12']=f'Chạy Postman CLI trên local Docker; {len(executed)} ca có assertion, các ca còn lại giữ Pending. Workbook đang liệt kê {count} Test Case ID (khác mô tả 152).'
wb['Cover']['F12']='Postman CLI 1.80; log trong workspace/postman/runs/collection-run.log'
stat=wb['Test Statistics']
stat['C6']=f'Round 1 đã chạy Postman CLI 1.80 trên local Docker Gateway. Có {len(executed)} ca có kết quả assertion; các ca chưa chạy giữ Pending. Danh sách hiện có {count} Test Case ID dù tài liệu ghi 152. Log: workspace/postman/runs/collection-run.log. Ca SRCH-08 thất bại: size=100000 trả HTTP 200 thay vì 400.'
# Excel should recalculate the workbook's existing COUNTIF summaries on open.
try:
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True
    wb.calculation.calcMode='auto'
except Exception: pass
wb.save(out)
print(f'Updated {out}: {count} cases; statuses: {len(executed)-1} Passed, 1 Failed, {count-len(executed)} Pending.')
